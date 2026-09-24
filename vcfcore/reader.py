"""Turn an uploaded CSV or XLSX file into rows of simple cell values.

Every cell comes back as a str, a date, a MonthDay (a date Excel shows without a
year), or None. Nothing here knows what the columns mean; that is columns.py.
"""

import csv
import io
import zipfile
from datetime import date, datetime

from openpyxl import load_workbook

from .models import MonthDay, TableError

MAX_ROWS = 2000

Cell = str | date | MonthDay | None
Rows = list[tuple[int, list[Cell]]]  # (spreadsheet row number, cells)


def read_table(stream, filename: str) -> Rows:
    """Read the first sheet of an .xlsx file, or a .csv file, into numbered rows."""
    data = stream.read()
    if not data:
        raise TableError("The file is empty.")

    name = filename.lower()
    if name.endswith(".xlsx"):
        rows = _read_xlsx(data)
    elif name.endswith(".csv"):
        rows = _read_csv(data)
    elif name.endswith(".xls"):
        raise TableError("Old .xls files are not supported. In Excel, use Save As and choose .xlsx or .csv.")
    else:
        raise TableError("Please upload a .csv, .xlsx, or .vcf file.")

    rows = [(number, cells) for number, cells in rows if any(_filled(c) for c in cells)]
    if not rows:
        raise TableError("The file has no data in it.")
    if len(rows) > MAX_ROWS + 1:  # + 1 for the header row
        raise TableError(f"The file has more than {MAX_ROWS} rows. Please split it into smaller files.")
    return rows


def _filled(cell: Cell) -> bool:
    return cell is not None and not (isinstance(cell, str) and not cell.strip())


def _read_xlsx(data: bytes) -> Rows:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        raise TableError("This .xlsx file could not be opened. It may be corrupted or password protected.") from exc

    try:
        sheet = workbook.worksheets[0]
        rows = []
        for number, row in enumerate(sheet.iter_rows(), start=1):
            rows.append((number, [_xlsx_cell(cell) for cell in row]))
            if len(rows) > MAX_ROWS * 2:  # stop early on huge sheets instead of reading them fully
                break
        return rows
    finally:
        workbook.close()


def _xlsx_cell(cell) -> Cell:
    value = getattr(cell, "value", None)
    if value is None:
        return None
    if isinstance(value, datetime):
        # typing "15-Mar" into Excel stores a full date with the current year and
        # formats it as "d-mmm"; if the format hides the year, the year is not real
        if "y" not in (cell.number_format or "").lower():
            return MonthDay(value.month, value.day)
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # a phone stored as a number, 2125550123.0 -> "2125550123"
    return str(value)


def _read_csv(data: bytes) -> Rows:
    text = decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel  # plain comma separated
    reader = csv.reader(io.StringIO(text), dialect)
    rows = []
    for number, row in enumerate(reader, start=1):
        rows.append((number, [cell for cell in row]))
        if len(rows) > MAX_ROWS * 2:
            break
    return rows


def decode(data: bytes) -> str:
    # utf-8-sig also strips the invisible marker Excel puts at the start of "CSV UTF-8" files
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")  # never fails, every byte maps to a character
